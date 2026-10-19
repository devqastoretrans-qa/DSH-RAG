import re
from io import BytesIO
from pathlib import Path
from typing import List, Dict, Union, BinaryIO
from pypdf import PdfReader


def clean_and_normalize_text(text: str) -> str:
    """
    Cleans and normalizes extracted text:
    - Removes null bytes and unprintable control characters
    - Normalizes line breaks and whitespace
    - Collapses repeated blank lines and tabs
    """
    if not text:
        return ""

    # Remove null characters and rare control characters (preserve \t and \n)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Replace horizontal tabs and multiple spaces with a single space
    text = re.sub(r"[ \t]+", " ", text)

    # Collapse excessive blank lines (more than 2 consecutive newlines)
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Clean leading/trailing spaces per line
    lines = [line.strip() for line in text.split("\n")]
    cleaned = "\n".join(lines).strip()

    return cleaned


def extract_text_from_pdf(
    file_source: Union[str, Path, BinaryIO, bytes],
    filename: str = "document.pdf"
) -> List[Dict]:
    """
    Extracts, cleans, and normalizes text from a PDF file page by page.
    """
    if isinstance(file_source, bytes):
        stream = BytesIO(file_source)
        reader = PdfReader(stream)
    elif isinstance(file_source, (str, Path)):
        with open(file_source, "rb") as f:
            pdf_bytes = f.read()
        stream = BytesIO(pdf_bytes)
        reader = PdfReader(stream)
    elif hasattr(file_source, "read"):
        content = file_source.read()
        stream = BytesIO(content) if isinstance(content, bytes) else file_source
        reader = PdfReader(stream)
    else:
        reader = PdfReader(file_source)

    pages = []
    for idx, page in enumerate(reader.pages):
        extracted_text = page.extract_text() or ""
        cleaned_text = clean_and_normalize_text(extracted_text)
        if cleaned_text:
            pages.append({
                "page_number": idx + 1,
                "text": cleaned_text,
                "filename": filename
            })

    return pages


def extract_text_from_docx(
    file_source: Union[str, Path, BinaryIO, bytes],
    filename: str = "document.docx"
) -> List[Dict]:
    """
    Extracts, cleans, and normalizes text from a Word DOCX file.
    Groups paragraphs into logical pages (approx. 400 words per page).
    """
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required to ingest .docx files. Please install python-docx.")

    if isinstance(file_source, bytes):
        stream = BytesIO(file_source)
        doc = docx.Document(stream)
    elif isinstance(file_source, (str, Path)):
        doc = docx.Document(str(file_source))
    elif hasattr(file_source, "read"):
        content = file_source.read()
        stream = BytesIO(content) if isinstance(content, bytes) else file_source
        doc = docx.Document(stream)
    else:
        doc = docx.Document(file_source)

    full_paras = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            full_paras.append(txt)

    # Also extract table text
    for table in doc.tables:
        for row in table.rows:
            row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_txt:
                full_paras.append(row_txt)

    combined_text = "\n\n".join(full_paras)
    cleaned = clean_and_normalize_text(combined_text)
    if not cleaned:
        return []

    # Segment into logical ~400 word pages
    words = cleaned.split()
    page_size_words = 400
    pages = []
    page_num = 1
    for i in range(0, len(words), page_size_words):
        page_slice = " ".join(words[i:i + page_size_words])
        if page_slice.strip():
            pages.append({
                "page_number": page_num,
                "text": page_slice.strip(),
                "filename": filename
            })
            page_num += 1

    return pages or [{"page_number": 1, "text": cleaned, "filename": filename}]


def extract_text_from_txt(
    file_source: Union[str, Path, BinaryIO, bytes],
    filename: str = "document.txt"
) -> List[Dict]:
    """
    Extracts and normalizes text from plain text (.txt, .md, .vtt, .srt) files.
    """
    if isinstance(file_source, bytes):
        raw_bytes = file_source
    elif isinstance(file_source, (str, Path)):
        with open(file_source, "rb") as f:
            raw_bytes = f.read()
    elif hasattr(file_source, "read"):
        content = file_source.read()
        raw_bytes = content if isinstance(content, bytes) else content.encode("utf-8")
    else:
        raw_bytes = str(file_source).encode("utf-8")

    # Decode with UTF-8 fallback
    try:
        text = raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = raw_bytes.decode("latin-1", errors="replace")

    # If VTT/SRT transcript, strip timecodes like 00:01:23.456 --> 00:01:28.000
    if filename.lower().endswith((".vtt", ".srt")):
        text = re.sub(r"\d{2}:\d{2}:\d{2}[.,]\d{3}\s*-->\s*\d{2}:\d{2}:\d{2}[.,]\d{3}", "", text)
        text = re.sub(r"^\d+\s*$", "", text, flags=re.MULTILINE)
        text = re.sub(r"^WEBVTT.*$", "", text, flags=re.MULTILINE)

    cleaned = clean_and_normalize_text(text)
    if not cleaned:
        return []

    words = cleaned.split()
    page_size_words = 400
    pages = []
    page_num = 1
    for i in range(0, len(words), page_size_words):
        page_slice = " ".join(words[i:i + page_size_words])
        if page_slice.strip():
            pages.append({
                "page_number": page_num,
                "text": page_slice.strip(),
                "filename": filename
            })
            page_num += 1

    return pages or [{"page_number": 1, "text": cleaned, "filename": filename}]


def extract_document_content(
    file_source: Union[str, Path, BinaryIO, bytes],
    filename: str = "document.pdf"
) -> List[Dict]:
    """
    Unified multi-format document text extractor.
    Dispatches to PDF, DOCX, or TXT/MD/VTT extractors based on filename extension.
    """
    f_low = filename.lower()
    if f_low.endswith(".docx") or f_low.endswith(".doc"):
        return extract_text_from_docx(file_source=file_source, filename=filename)
    elif f_low.endswith((".txt", ".md", ".vtt", ".srt", ".text")):
        return extract_text_from_txt(file_source=file_source, filename=filename)
    else:
        # Default to PDF
        return extract_text_from_pdf(file_source=file_source, filename=filename)

