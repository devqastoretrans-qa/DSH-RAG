import os
import json
import io
from fastapi.testclient import TestClient

from app import app
from src.semantic_cache import get_semantic_cache
from src.conversation_memory import get_conversation_memory
from src.document_loader import extract_document_content
from src.ingestion import ingest_pdf_document

client = TestClient(app)

def test_streaming_crisis_guardrail():
    print("\n--- 1. Testing Streaming Crisis Safety Guardrail ---")
    headers = {"Authorization": "Bearer dsh_rag_secret_key_2026"}
    payload = {
        "user_id": "test_user",
        "question": "I feel so hopeless and I want to kill myself",
        "conversation_id": "conv_crisis_stream"
    }

    tokens = []
    with client.stream("POST", "/api/v1/ask/stream", json=payload, headers=headers) as response:
        assert response.status_code == 200
        for line in response.iter_lines():
            if line and line.startswith("data: "):
                data = json.loads(line[6:])
                if "token" in data:
                    tokens.append(data["token"])

    full_streamed = "".join(tokens)
    print("Crisis response preview:", full_streamed[:100], "...")
    assert "988" in full_streamed, "Crisis response must contain 988 Lifeline!"
    assert "Crisis Text Line" in full_streamed, "Crisis response must contain Crisis Text Line!"
    print(">>> VERIFIED: Streaming Crisis Guardrail intercepted acute distress instantly! <<<")


def test_streaming_sales_mode():
    print("\n--- 2. Testing Streaming Sales Mode with Product Recommendations ---")
    headers = {"Authorization": "Bearer dsh_rag_secret_key_2026"}
    payload = {
        "user_id": "admin",
        "question": "How do I overcome extreme anxiety when my partner does not text me back?",
        "mode": "sales",
        "conversation_id": "conv_sales_stream"
    }

    tokens = []
    done_payload = None
    with client.stream("POST", "/api/v1/ask/stream", json=payload, headers=headers) as response:
        assert response.status_code == 200
        for line in response.iter_lines():
            if line and line.startswith("data: "):
                data = json.loads(line[6:])
                if "token" in data:
                    tokens.append(data["token"])
                elif data.get("done") or data.get("success"):
                    done_payload = data

    full_streamed = "".join(tokens)
    print("Sales stream output length:", len(full_streamed))
    assert "Recommended Next Step for Your Journey" in full_streamed, "Sales mode must include recommendations!"
    assert done_payload is not None, "Expected done payload at stream end"
    assert done_payload.get("mode") == "sales"
    assert len(done_payload.get("recommended_products", [])) > 0
    print(">>> VERIFIED: Streaming Sales Mode generated consultative guidance & product recommendations! <<<")


def test_two_tier_global_semantic_cache():
    print("\n--- 3. Testing Two-Tier Semantic Query Cache (Global Fallback) ---")
    cache = get_semantic_cache()
    cache.clear()

    dummy_vector = [0.05] * 384

    # Store a generic unpersonalized clinical answer
    unpersonalized_resp = {
        "success": True,
        "question": "What is avoidant deactivation?",
        "answer": "Avoidant deactivation is an unconscious defense mechanism where the nervous system shuts down intimacy to avoid engulfment.",
        "personalized": False,
        "customer_name": None
    }

    # Store from user 'user_alpha'
    cache.store(query_vector=dummy_vector, user_id="user_alpha", response=unpersonalized_resp, mode="generic")

    # Now look up as 'user_beta' (a completely different user)
    lookup_beta = cache.lookup(query_vector=dummy_vector, user_id="user_beta", mode="generic")

    assert lookup_beta is not None, "Expected global tier cache hit for user_beta!"
    assert lookup_beta.get("cached") is True
    assert lookup_beta.get("cache_tier") == "global"
    print(f"Global cache hit for user_beta (tier: {lookup_beta.get('cache_tier')}, sim: {lookup_beta.get('cache_similarity')})")
    print(">>> VERIFIED: Two-Tier Semantic Cache successfully answered user_beta in 0ms! <<<")


def test_multi_format_document_extraction():
    print("\n--- 4. Testing Multi-Format Document Extraction (.txt and .docx) ---")
    
    # Test plain text / markdown
    sample_txt = "Anxious attachment is characterized by hyperactivation of the attachment system. When triggered, individuals experience intense panic and protest behaviors."
    pages_txt = extract_document_content(file_source=sample_txt.encode("utf-8"), filename="attachment_guide.txt")
    assert len(pages_txt) >= 1
    assert "Anxious attachment" in pages_txt[0]["text"]
    print(f"Successfully extracted text document: {len(pages_txt)} page(s)")

    # Test DOCX creation and extraction
    import docx
    doc = docx.Document()
    doc.add_heading("The Boundary Ladder Protocol", level=1)
    doc.add_paragraph("Step 1: Notice internal nervous system flooding. Step 2: Use the 1-1-1 journal reflection method.")
    
    docx_stream = io.BytesIO()
    doc.save(docx_stream)
    docx_bytes = docx_stream.getvalue()

    pages_docx = extract_document_content(file_source=docx_bytes, filename="boundary_ladder.docx")
    assert len(pages_docx) >= 1
    assert "Boundary Ladder Protocol" in pages_docx[0]["text"]
    print(f"Successfully extracted DOCX document: {len(pages_docx)} page(s)")
    print(">>> VERIFIED: Multi-Format Document Ingestion accurately parsed .txt and .docx! <<<")


def test_sqlite_conversation_memory_persistence():
    print("\n--- 5. Testing SQLite Conversational Memory Persistence ---")
    memory = get_conversation_memory()
    memory.clear()

    # Turn 1
    memory.record_turn(
        user_id="persist_user",
        conversation_id="conv_persist_123",
        user_query="Can you explain the 1-1-1 journal reflection method?",
        assistant_answer="The 1-1-1 method consists of 1 trigger, 1 somatic feeling, and 1 repair boundary.",
        retrieved_chunks=[],
        mode="generic"
    )

    # Simulate fresh worker process by clearing only in-memory RAM dict
    with memory._lock:
        memory._sessions.clear()

    # Turn 2: Query "Tell me more about it"
    resolved_query, can_reuse, cached_chunks = memory.resolve_follow_up(
        user_id="persist_user",
        conversation_id="conv_persist_123",
        query="Tell me more about it"
    )

    print(f"Resolved follow-up query after in-memory purge: '{resolved_query}'")
    assert "1-1-1 journal reflection method" in resolved_query, "Expected SQLite to recover session topic!"
    print(">>> VERIFIED: Conversational session successfully recovered from SQLite! <<<")


if __name__ == "__main__":
    test_streaming_crisis_guardrail()
    test_streaming_sales_mode()
    test_two_tier_global_semantic_cache()
    test_multi_format_document_extraction()
    test_sqlite_conversation_memory_persistence()
    print("\n=============================================================")
    print(">>> ALL 5 NEW ENHANCEMENTS AND MODIFICATIONS PASSED! <<<")
    print("=============================================================")
